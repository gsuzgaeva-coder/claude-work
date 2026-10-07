"""
Выгрузка фактических приёмок Яндекс Маркета (поставки FBY) по всем кабинетам из .env.

Запуск:  python scripts/ym_supplies_export.py
Результат: out/ЯМ_приёмки_<дата>.xlsx и out/ym_priemki_lines.csv
Правила отбора заявок — knowledge/02_яндекс_маркет_поставки.md
Артикул и размер берутся из data/spravochnik_tovarov.csv по баркоду (не угадываются по артикулу кабинета).
"""
from __future__ import annotations

import csv
import json
import os
import sys
import time
from collections import Counter, OrderedDict
from datetime import date
from pathlib import Path

import httpx
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
API = "https://api.partner.market.yandex.ru"
CACHE = ROOT / "data" / "cache" / "ym_requests.json"
TOTALS = ROOT / "data" / "cache" / "ym_last_totals.json"
OVERRIDES = ROOT / "data" / "barcode_overrides.csv"
SPRAV = ROOT / "data" / "spravochnik_tovarov.csv"
OUT = ROOT / "out"

CHILD = "VIRTUAL_DISTRIBUTION_CENTER_CHILD"
PARENT = "VIRTUAL_DISTRIBUTION_CENTER"
STATUS_RU = {
    "WAREHOUSE_HANDLING": "На складе, идёт приёмка",
    "ARRIVED_TO_SERVICE": "Прибыла на склад",
    "SHIPPED_TO_SERVICE": "В пути на склад",
    "ACCEPTED_BY_WAREHOUSE_SYSTEM": "Создана, ещё не прибыла",
}


def load_env() -> None:
    env = ROOT / ".env"
    if not env.exists():
        sys.exit("Нет файла .env — скопируйте .env.example в .env и впишите ключи кабинетов.")
    for line in env.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip().strip('"').strip("'"))


def cabinets() -> list[dict]:
    out = []
    for code in [c.strip() for c in os.environ.get("YM_CABINETS", "").split(",") if c.strip()]:
        cab = {f: os.environ.get(f"YM_{code}_{f.upper()}", "") for f in ("token", "business_id", "campaign_id", "ip")}
        missing = [f for f, v in cab.items() if not v]
        if missing:
            sys.exit(f"В .env для кабинета {code} не заполнено: {', '.join('YM_' + code + '_' + m.upper() for m in missing)}")
        cab["code"] = code
        cab["ip_name"] = os.environ.get(f"YM_{code}_IP_NAME") or cab["ip"]
        out.append(cab)
    if not out:
        sys.exit("В .env не задан YM_CABINETS (например: YM_CABINETS=VIDINEEVA,POZOYAN).")
    return out


def call(client: httpx.Client, token: str, path: str, body: dict | None, params: dict) -> dict:
    for attempt in range(6):
        try:
            r = client.post(API + path, headers={"Api-Key": token}, json=body or {}, params=params)
        except httpx.HTTPError as exc:
            if attempt == 5:
                raise
            print(f"  сеть: {exc!r}, повтор", flush=True)
            time.sleep(10)
            continue
        if r.status_code in (420, 429) or r.status_code >= 500:
            time.sleep(15 * (attempt + 1))
            continue
        if r.status_code == 401:
            sys.exit("Маркет ответил 401: ключ недействителен или отозван. Проверьте .env.")
        if r.status_code == 403:
            sys.exit("Маркет ответил 403: у ключа нет прав на поставки/каталог.")
        if r.status_code >= 400:
            sys.exit(f"{path} → {r.status_code}: {r.text[:300]}")
        return r.json()
    sys.exit(f"{path}: Маркет не отвечает после нескольких попыток, запустите позже.")


def pages(client, token, path, key, body=None, limit=100) -> list[dict]:
    out, tok = [], None
    for _ in range(200):
        params = {"limit": limit}
        if tok:
            params["page_token"] = tok
        res = call(client, token, path, body, params).get("result") or {}
        out += res.get(key) or []
        tok = (res.get("paging") or {}).get("nextPageToken")
        if not tok:
            break
    return out


def warehouse(name: str) -> str:
    for needle, short in (("Софьино Транзит", "Софьино Х-док"), ("Софьино", "Софьино"), ("Ростов", "Ростов-на-Дону"),
                          ("Екатеринбург", "Екатеринбург"), ("Казань", "Казань"), ("Шушары", "Шушары"), ("Самара", "Самара")):
        if needle in name:
            return short
    return name.strip()


def ru(d: str) -> str:
    return f"{d[8:10]}.{d[5:7]}.{d[:4]}" if d else ""


def collect(cab: dict, cache: dict) -> tuple[list[dict], dict]:
    """Вернуть заявки кабинета с составами (с кэшем по updatedAt) и карту offerId → баркоды."""
    old = cache.get(cab["code"], {})
    with httpx.Client(timeout=httpx.Timeout(60, connect=15)) as client:
        reqs = pages(client, cab["token"], f"/v2/campaigns/{cab['campaign_id']}/supply-requests", "requests",
                     body={"sorting": {"direction": "DESC", "attribute": "UPDATED_AT"}})
        print(f"{cab['code']}: заявок {len(reqs)}", flush=True)
        fresh, done = {}, 0
        for r in reqs:
            if r.get("type") != "SUPPLY" or r.get("subtype") != CHILD:
                continue
            rid = str(r["id"]["id"])
            prev = old.get(rid)
            if prev and prev["request"].get("updatedAt") == r.get("updatedAt"):
                fresh[rid] = prev
                continue
            items = pages(client, cab["token"], f"/v2/campaigns/{cab['campaign_id']}/supply-requests/items", "items",
                          body={"requestId": r["id"]["id"]})
            fresh[rid] = {"request": r, "items": items}
            done += 1
            if done % 25 == 0:
                print(f"  {cab['code']}: прочитано составов {done}", flush=True)
        print(f"{cab['code']}: поставок с ФФ {len(fresh)}, перечитано {done}", flush=True)
        offers = {}
        for om in pages(client, cab["token"], f"/v2/businesses/{cab['business_id']}/offer-mappings", "offerMappings", limit=200):
            o = om.get("offer") or {}
            offers[o.get("offerId")] = o.get("barcodes") or []
    cache[cab["code"]] = fresh
    return list(fresh.values()), offers


def build(cabs: list[dict], cache: dict):
    overrides, sprav = {}, {}
    if OVERRIDES.exists():
        for row in csv.DictReader(OVERRIDES.open(encoding="utf-8")):
            overrides[row["offer_id"]] = row
    if SPRAV.exists():
        for row in csv.DictReader(SPRAV.open(encoding="utf-8")):
            sprav[row["barcode"]] = row
    lines, requests, no_barcode, no_sprav = [], [], Counter(), Counter()
    for cab in cabs:
        entries, offers = collect(cab, cache)
        for e in entries:
            r = e["request"]
            loc = r.get("targetLocation") or {}
            c = r.get("counters") or {}
            head = {
                "ip": cab["ip"], "date": (loc.get("requestedDate") or "")[:10],
                "mpid": str(r["id"].get("marketplaceRequestId") or r["id"]["id"]),
                "vrc": f"ВРЦ-{((r.get('parentLink') or {}).get('id') or {}).get('id', '')}",
                "wh": warehouse(loc.get("name") or ""), "status": r["status"],
            }
            requests.append({**head, "plan": c.get("planCount") or 0, "fact": c.get("factCount"),
                             "short": c.get("shortageCount") or 0, "surplus": c.get("surplusCount") or 0,
                             "unacc": c.get("unacceptableCount") or 0})
            finished = r["status"] == "FINISHED"
            for it in e["items"]:
                offer = it["offerId"]
                ic = it.get("counters") or {}
                bcs = offers.get(offer) or []
                barcode = next((b for b in bcs if b.startswith("20")), bcs[0] if bcs else "")
                note = ""
                if not barcode and offer in overrides:
                    barcode = overrides[offer]["barcode"]
                    note = overrides[offer].get("note", "")
                if not barcode and finished:
                    no_barcode[(cab["ip"], offer)] += ic.get("factCount") or 0
                ref = sprav.get(barcode) or {}
                if barcode and not ref and finished:
                    no_sprav[(cab["ip"], barcode, offer)] += ic.get("factCount") or 0
                lines.append({**head, "offer": offer, "barcode": barcode, "note": note,
                              "art": ref.get("artikul", ""), "size": ref.get("razmer", ""),
                              "plan": ic.get("planCount") or 0,
                              "fact": (ic.get("factCount") or 0) if finished else None,
                              "short": ic.get("shortageCount") or 0, "surplus": ic.get("surplusCount") or 0})
    return lines, requests, no_barcode, no_sprav


def write(lines, requests, cabs) -> Path:
    OUT.mkdir(exist_ok=True)
    fin = sorted((l for l in lines if l["status"] == "FINISHED"), key=lambda l: (l["date"], l["ip"], l["mpid"], l["offer"]))
    fin_req = sorted((r for r in requests if r["status"] == "FINISHED"), key=lambda r: (r["date"], r["ip"], r["mpid"]))
    open_req = sorted((r for r in requests if r["status"] not in ("FINISHED", "CANCELLED")), key=lambda r: (r["ip"], r["date"]))
    bold, fill = Font(bold=True), PatternFill("solid", fgColor="DDEBF7")
    wb = Workbook()
    today = f"{date.today():%d.%m.%Y}"

    def sheet(title, header, rows, widths):
        ws = wb.create_sheet(title)
        ws.append(header)
        for cell in ws[1]:
            cell.font, cell.fill, cell.alignment = bold, fill, Alignment(wrap_text=True, vertical="center")
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def totals(label, year, part):
        return [label, year, len({l["mpid"] for l in part}), sum(l["plan"] for l in part), sum(l["fact"] for l in part),
                sum(l["short"] for l in part), sum(l["surplus"] for l in part)]

    first = wb.active
    first.title = "Читать сначала"
    starts = ", ".join(f"с {ru(min(l['date'] for l in fin if l['ip'] == c['ip']))} ({c['ip_name']})"
                       for c in cabs if any(l["ip"] == c["ip"] for l in fin))
    for text in [
        "ПРИЁМКИ НА ЯНДЕКС МАРКЕТЕ — фактические данные для учёта", None,
        f"Что это: все поставки с фулфилментов на склады Яндекс Маркета (FBY), которые Маркет принял. "
        f"Источник — кабинеты продавца ЯМ (API), выгрузка {today}.",
        f"Период: {starts} по {today}. Всего принято {sum(l['fact'] for l in fin)} шт по {len(fin_req)} поставкам.", None,
        "ЛИСТЫ",
        "• Свод — итоги по ИП и годам.",
        "• Приёмки по заявкам — одна строка = одна поставка: дата, номер заявки ЯМ, склад, заявлено / принято / недостача / излишек.",
        "• Приёмки построчно — одна строка = товар в поставке (баркод, артикул, размер, заявлено, принято). "
        "Это основная таблица для загрузки в учёт.",
        "• По баркодам — сколько принято каждого баркода по годам и всего.",
        "• Ещё не принято — поставки, которые отгружены или созданы, но Маркет их пока не принял. В учёт приёмок не входят.", None,
        "КАК ЧИТАТЬ",
        "• «Принято» — фактически принято складом Маркета. «Заявлено» — количество в заявке. Недостача = заявлено − принято.",
        "• Дата — плановая дата приёмки по заявке ЯМ (дату отгрузки с фулфилмента Маркет не хранит). "
        "Для транзита это дата на конечном складе.",
        "• «ВРЦ-…» — номер родительской (транзитной) заявки; «№ заявки ЯМ» — номер поставки на конкретный склад, "
        "по нему заявка ищется в кабинете.",
        "• Артикул наш и размер — из справочника товаров по баркоду. Пусто — баркода нет в справочнике, его надо добавить.", None,
        "ЧТО НЕ ВХОДИТ",
        "• Отменённые заявки (если товар перевыставили новой заявкой — он учтён в новой).",
        "• Перемещения между складами Маркета и находки инвентаризации — это не отгрузки с фулфилмента.",
        "• Возвраты покупателей, вывозы брака и непринятого со складов Маркета.", None,
        "ОГОВОРКИ",
        "• Строки с текстом в колонке «Примечание» — баркод подставлен вручную, его нужно подтвердить.",
    ]:
        first.append([text])
    first["A1"].font = Font(bold=True, size=13)
    first.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Свод")
    ws.append(["ИП", "Год", "Принятых поставок", "Заявлено, шт", "Принято, шт", "Недостача, шт", "Излишек, шт"])
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    for cab in cabs:
        own = [l for l in fin if l["ip"] == cab["ip"]]
        for year in sorted({l["date"][:4] for l in own}):
            ws.append(totals(cab["ip_name"], year, [l for l in own if l["date"][:4] == year]))
        ws.append(totals(cab["ip_name"], "Итого", own))
        for cell in ws[ws.max_row]:
            cell.font = bold
    ws.append(totals("Все ИП", "Итого", fin))
    for cell in ws[ws.max_row]:
        cell.font = bold
    for i, w in enumerate([32, 8, 12, 12, 12, 12, 12], 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    sheet("Приёмки по заявкам",
          ["Дата приёмки", "ИП", "№ заявки ЯМ", "ВРЦ (транзитная заявка)", "Склад ЯМ", "Заявлено, шт", "Принято, шт",
           "Недостача, шт", "Излишек, шт", "Не принято (аномалии), шт"],
          [[ru(r["date"]), r["ip"], r["mpid"], r["vrc"], r["wh"], r["plan"], r["fact"], r["short"], r["surplus"], r["unacc"]]
           for r in fin_req], [13, 7, 13, 22, 18, 12, 12, 13, 12, 16])
    header = ["Дата приёмки", "ИП", "№ заявки ЯМ", "ВРЦ (транзитная заявка)", "Склад ЯМ", "Баркод", "Артикул наш", "Размер",
              "Артикул в кабинете ЯМ", "Заявлено, шт", "Принято, шт", "Недостача, шт", "Излишек, шт", "Примечание"]
    rows = [[ru(l["date"]), l["ip"], l["mpid"], l["vrc"], l["wh"], l["barcode"], l["art"], l["size"], l["offer"], l["plan"],
             l["fact"], l["short"], l["surplus"], l["note"]] for l in fin]
    sheet("Приёмки построчно", header, rows, [13, 7, 13, 22, 18, 16, 30, 8, 36, 11, 11, 12, 11, 36])
    agg = OrderedDict()
    for l in sorted(fin, key=lambda l: (l["ip"], l["art"] or l["offer"], l["barcode"])):
        a = agg.setdefault((l["ip"], l["barcode"]), {"art": l["art"], "size": l["size"], "n": Counter()})
        a["n"][l["date"][:4]] += l["fact"]
        a["n"]["all"] += l["fact"]
        a["n"]["plan"] += l["plan"]
    years = sorted({l["date"][:4] for l in fin})
    sheet("По баркодам", ["ИП", "Баркод", "Артикул наш", "Размер"] + [f"Принято {y}, шт" for y in years]
          + ["Принято всего, шт", "Заявлено всего, шт"],
          [[k[0], k[1], v["art"], v["size"]] + [v["n"][y] for y in years] + [v["n"]["all"], v["n"]["plan"]] for k, v in agg.items()],
          [7, 16, 30, 8] + [14] * (len(years) + 2))
    sheet("Ещё не принято", ["ИП", "Плановая дата приёмки", "№ заявки ЯМ", "ВРЦ (транзитная заявка)", "Склад ЯМ",
                             f"Статус на {today}", "Заявлено, шт"],
          [[r["ip"], ru(r["date"]), r["mpid"], r["vrc"], r["wh"], STATUS_RU.get(r["status"], r["status"]), r["plan"]]
           for r in open_req], [7, 14, 13, 22, 18, 28, 12])
    path = OUT / f"ЯМ_приёмки_{date.today():%d-%m-%Y}.xlsx"
    wb.save(path)
    with (OUT / "ym_priemki_lines.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(header)
        w.writerows(rows)
    return path


def main() -> None:
    load_env()
    cabs = cabinets()
    cache = json.loads(CACHE.read_text(encoding="utf-8")) if CACHE.exists() else {}
    lines, requests, no_barcode, no_sprav = build(cabs, cache)
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    CACHE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")

    fin = [l for l in lines if l["status"] == "FINISHED"]
    by_req = Counter()
    for l in fin:
        by_req[l["mpid"]] += l["fact"]
    mismatch = [(r["mpid"], r["fact"], by_req[r["mpid"]]) for r in requests
                if r["status"] == "FINISHED" and (r["fact"] or 0) != by_req[r["mpid"]]]
    path = write(lines, requests, cabs)

    prev = json.loads(TOTALS.read_text(encoding="utf-8")) if TOTALS.exists() else {}
    now = {cab["ip"]: sum(l["fact"] for l in fin if l["ip"] == cab["ip"]) for cab in cabs}
    dropped = {ip: (prev[ip], n) for ip, n in now.items() if ip in prev and n < prev[ip]}
    if not dropped:
        TOTALS.write_text(json.dumps(now, ensure_ascii=False), encoding="utf-8")

    print("\n=== ИТОГ ===")
    for cab in cabs:
        part = [l for l in fin if l["ip"] == cab["ip"]]
        print(f"{cab['ip']}: принятых поставок {len({l['mpid'] for l in part})}, заявлено {sum(l['plan'] for l in part)}, "
              f"принято {sum(l['fact'] for l in part)}")
    if dropped:
        print("ВНИМАНИЕ: «принято всего» уменьшилось с прошлой выгрузки — в журнал НЕ загружать, разобраться:")
        for ip, (was, n) in dropped.items():
            print(f"  {ip}: было {was}, стало {n}")
    waiting = [r for r in requests if r["status"] not in ("FINISHED", "CANCELLED")]
    print(f"Ещё не принято: поставок {len(waiting)}, заявлено {sum(r['plan'] for r in waiting)} шт")
    print("Сверка суммы строк с итогами заявок:", "сошлось" if not mismatch else f"РАСХОЖДЕНИЯ: {mismatch}")
    if no_barcode:
        print("Строки без баркода (добавьте в data/barcode_overrides.csv):")
        for (ip, offer), qty in no_barcode.items():
            print(f"  {ip}  {offer}  принято {qty} шт")
    if no_sprav:
        print("Баркоды не из справочника товаров (добавьте в data/spravochnik_tovarov.csv):")
        for (ip, bc, offer), qty in no_sprav.items():
            print(f"  {ip}  {bc}  {offer}  принято {qty} шт")
    print("Файл:", path)


if __name__ == "__main__":
    main()
